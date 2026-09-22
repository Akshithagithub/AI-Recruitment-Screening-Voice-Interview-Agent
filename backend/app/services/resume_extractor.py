from pathlib import Path

import fitz
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


class ResumeExtractionError(Exception):
    """Raised when resume text extraction fails."""


def extract_pdf_text(file_path: Path) -> str:
    """
    Extract text from a PDF resume using PyMuPDF.
    """

    try:
        document = fitz.open(file_path)

        pages = []

        for page in document:
            text = page.get_text("text")

            if text:
                pages.append(text)

        document.close()

        extracted_text = "\n".join(pages).strip()

        if not extracted_text:
            raise ResumeExtractionError(
                "No text could be extracted from the PDF"
            )

        return extracted_text

    except ResumeExtractionError:
        raise

    except Exception as error:
        raise ResumeExtractionError(
            "Failed to extract text from PDF"
        ) from error


def extract_docx_text(file_path: Path) -> str:
    """
    Extract text from a DOCX resume using python-docx.
    """

    try:
        document = Document(file_path)

        paragraphs = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        # Also extract text from tables.
        for table in document.tables:
            for row in table.rows:
                row_text = []

                for cell in row.cells:
                    text = cell.text.strip()

                    if text:
                        row_text.append(text)

                if row_text:
                    paragraphs.append(" | ".join(row_text))

        extracted_text = "\n".join(paragraphs).strip()

        if not extracted_text:
            raise ResumeExtractionError(
                "No text could be extracted from the DOCX"
            )

        return extracted_text

    except ResumeExtractionError:
        raise

    except Exception as error:
        raise ResumeExtractionError(
            "Failed to extract text from DOCX"
        ) from error


def extract_resume_text(file_path: str | Path) -> str:
    """
    Extract text from a PDF or DOCX resume.
    """

    path = Path(file_path)

    if not path.exists():
        raise ResumeExtractionError(
            "Resume file not found"
        )

    extension = path.suffix.lower()

    if extension == ".pdf":
        return extract_pdf_text(path)

    if extension == ".docx":
        return extract_docx_text(path)

    raise ResumeExtractionError(
        f"Unsupported resume format: {extension}"
    )